import pandas as pd
import random
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import string

# قائمة الموردين الوهميين
suppliers = ["شركة النيل الذهبية", "مؤسسة الخليج التجارية", "المصرية للتوزيع", "شركة المشرق", "الشركة العربية للتجارة"]

# قائمة المنتجات الغذائية (بدون خضار وفاكهة) + تنظيف + عطور
products_data = {
    "الحبوب والدقيق": ["أرز أبيض 1كج", "دقيق حنطة 1كج", "ذرة صفراء 1كج", "حنطة سوداء 500غ", "شعير 1كج"],
    "الزيوت والدهون": ["زيت الذرة 1لتر", "زيت الزيتون 750مل", "زيت الكانولا 1.5لتر", "السمن النباتي 500غ", "سمن حيواني 400غ"],
    "الألبان والأجبان": ["حليب طازج 1لتر", "لبن رايب 500مل", "جبن أبيض 200غ", "جبن أصفر 150غ", "زبادي طبيعي 400غ"],
    "اللحوم المعلبة": ["سمك معلب 170غ", "دجاج معلب 400غ", "لحم مفروم معلب 300غ", "سمك بالطماطم 150غ", "دجاج بالزيت 200غ"],
    "الحلويات والشوكولاتة": ["شوكولاتة داكنة 100غ", "حلوى تمر معلبة 200غ", "كيك شوكولاتة 150غ", "بسكويت زبدة 200غ", "ويفر بالشوكولاتة 100غ"],
    "الشاي والقهوة": ["شاي أسود 500غ", "قهوة مطحونة 250غ", "شاي أخضر 200غ", "قهوة سريعة الذوبان 100غ", "شاي البابونج 150غ"],
    "المشروبات": ["عصير برتقال 1لتر", "عصير تفاح 750مل", "ماء معدني 1.5لتر", "شراب ليموني 500مل", "عصير عنب 1لتر"],
    "الصلصات والتوابل": ["صلصة طماطم 400غ", "مايونيز 250غ", "خل أبيض 750مل", "صويا صوص 200مل", "عسل نحل 400غ"],
    "الحبوب الجاهزة": ["كورن فليكس 200غ", "جرانولا 300غ", "ميوسلي 250غ", "فشار للفرقعة 100غ", "شوفان سريع 200غ"],
    "المعكرونة والأرز": ["باستا إسباغيتي 500غ", "معكرونة ملتوية 250غ", "أرز بسمتي 1كج", "أرز مصري 1كج", "معكرونة حنطة 400غ"],
    "تنظيف الأرضيات": ["ديتول أرضي 750مل", "فينو أرضي 1لتر", "ملمع أرضي 500مل", "منظف حجري 750مل", "مزيل شمع 500مل"],
    "تنظيف الأطباق": ["سائل جلي 500مل", "جل تنظيف قوي 750مل", "إسفنجة تنظيف 3قطع", "سائل تنظيف 1لتر", "مسحوق غسيل 1كج"],
    "تنظيف الحمام": ["مطهر حمام 500مل", "تنظيف الدش 750مل", "ملمع البلاط 500مل", "إزالة الكلس 1لتر", "معطر حمام 300مل"],
    "مستحضرات التنظيف": ["مناديل مبللة 100قطعة", "فرشاة تنظيف 1قطعة", "قفاز تنظيف 5ثنائيات", "كيس قمامة 50قطعة", "ورق تنظيف 200ورقة"],
    "العطور والمعطرات": ["عطر جنى الزهور 50مل", "عطر العود الفاخر 100مل", "معطر الهواء 300مل", "عطر الورد 75مل", "معطر السيارة 10مل"],
    "معطرات الفم": ["غسول فم معقم 500مل", "معجون أسنان 100مل", "رذاذ تنفس 15مل", "عبادة فم 25مل", "معطر فم نعناع 50مل"],
}

# توسيع المنتجات لتصل إلى 150
all_products = []
for category, items in products_data.items():
    for item in items:
        all_products.append({"name": item, "category": category})

# إضافة المزيد من المنتجات للوصول إلى 150
additional_products = [
    {"name": f"منتج إضافي {i}", "category": list(products_data.keys())[i % len(products_data)]}
    for i in range(150 - len(all_products))
]
all_products.extend(additional_products)

def generate_barcode():
    """توليد باركود 12 رقم"""
    return ''.join([str(random.randint(0, 9)) for _ in range(12)])

def create_invoice(supplier_name, invoice_number):
    """إنشاء فاتورة واحدة"""

    # اختيار عدد عشوائي من المنتجات (من 10 إلى 30)
    num_items = random.randint(10, 30)
    selected_products = random.sample(all_products, min(num_items, len(all_products)))

    # إنشاء Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "الفاتورة"

    # تعيين عرض الأعمدة
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 20
    ws.column_dimensions['H'].width = 15
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 15
    ws.column_dimensions['K'].width = 20

    # تعريب القالب من اليمين لليسار
    ws.right_to_left = True

    # الرأس
    ws['A1'] = "الرقم"
    ws['B1'] = "الوصف"
    ws['C1'] = "الكمية"
    ws['D1'] = "الوحدة"
    ws['E1'] = "السعر"
    ws['F1'] = "صافي البيع"
    ws['G1'] = "الصلاحية"
    ws['H1'] = "المورية"
    ws['I1'] = "فرعي"
    ws['J1'] = "رئيسي"
    ws['K1'] = "نوع المنتج"

    # تنسيق الرأس
    header_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    header_font = Font(bold=True, size=11)

    for col in ['A1', 'B1', 'C1', 'D1', 'E1', 'F1', 'G1', 'H1', 'I1', 'J1', 'K1']:
        ws[col].fill = header_fill
        ws[col].font = header_font
        ws[col].alignment = Alignment(horizontal="center", vertical="center")

    # إضافة البيانات
    row = 2
    total = 0

    for product in selected_products:
        barcode = generate_barcode()
        quantity = random.randint(5, 100)
        unit = random.choice(["كج", "لتر", "علبة", "قطعة", "كيس", "صندوق"])
        price = round(random.uniform(5, 500), 2)
        net_price = round(price * 0.95, 2)  # خصم 5%

        # تاريخ الصلاحية (من 3 أشهر إلى سنة من الآن)
        days_ahead = random.randint(90, 365)
        expiry_date = (datetime.now() + timedelta(days=days_ahead)).strftime("%d/%m/%Y")

        amount = round(quantity * price, 2)
        total += amount

        ws[f'A{row}'] = barcode
        ws[f'B{row}'] = product['name']
        ws[f'C{row}'] = quantity
        ws[f'D{row}'] = unit
        ws[f'E{row}'] = price
        ws[f'F{row}'] = net_price
        ws[f'G{row}'] = expiry_date
        ws[f'H{row}'] = supplier_name
        ws[f'I{row}'] = amount
        ws[f'J{row}'] = round(amount * 1.05, 2)  # مبلغ رئيسي مع 5% إضافي
        ws[f'K{row}'] = product['category']

        # تنسيق الخلايا
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
            ws[f'{col}{row}'].alignment = Alignment(horizontal="center", vertical="center")

        row += 1

    # إضافة الإجمالي
    total_row = row + 1
    ws[f'H{total_row}'] = "الإجمالي:"
    ws[f'I{total_row}'] = total
    ws[f'J{total_row}'] = round(total * 1.05, 2)

    total_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    total_font = Font(bold=True)

    for col in ['H', 'I', 'J']:
        ws[f'{col}{total_row}'].fill = total_fill
        ws[f'{col}{total_row}'].font = total_font

    # حفظ الملف باسم المورد
    filename = f"{supplier_name.replace(' ', '_')}_invoice_{invoice_number}.xlsx"
    filepath = f"C:\\Users\\naji2\\DataLY_6\\{filename}"
    wb.save(filepath)

    print(f"✓ تم إنشاء: {filename}")
    return filepath

# إنشاء 5 فواتير
print("جاري إنشاء 5 فواتير...")
for i, supplier in enumerate(suppliers, 1):
    create_invoice(supplier, i)

print("\n✓ تم إنشاء جميع الفواتير بنجاح!")
